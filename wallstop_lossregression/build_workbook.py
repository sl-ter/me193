"""Build gradient_descent.xlsx: one formula row per gradient-descent epoch."""
import csv
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import LineChart, ScatterChart, Reference, Series
from openpyxl.comments import Comment

HERE = Path(__file__).parent

BLUE = Font(name="Arial", color="0000FF"); BOLD = Font(name="Arial", bold=True); BASE = Font(name="Arial")
YELLOW = PatternFill("solid", fgColor="FFFF00")

rows = list(csv.reader(open(HERE / "annacode-data.csv", encoding="utf-8")))[1:]
wall = sorted((float(r[1]), float(r[2])) for r in rows if float(r[1]) > 0 and float(r[2]) != 0)

def build(ws, X, Y, mode, lr, epochs, xname="x", yname="Y", note=None):
    n = len(X); last = 1 + n
    ws["A1"], ws["B1"] = xname, yname
    for i, (x, y) in enumerate(zip(X, Y), start=2):
        ws[f"A{i}"], ws[f"B{i}"] = x, y
        for c in "AB": ws[f"{c}{i}"].font = BLUE; ws[f"{c}{i}"].fill = YELLOW
    ar = f"$A$2:$A${last}"
    if mode == "center":
        ws["C1"] = "x centered"
        for i in range(2, last + 1): ws[f"C{i}"] = f"=A{i}-AVERAGE({ar})"
    elif mode == "standardize":
        ws["C1"] = "x standardized"
        for i in range(2, last + 1): ws[f"C{i}"] = f"=(A{i}-AVERAGE({ar}))/STDEVP({ar})"
    col = "A" if mode == "raw" else "C"
    xr, yr = f"${col}$2:${col}${last}", f"$B$2:$B${last}"

    ws["E1"], ws["F1"] = "learning rate", lr
    ws["E2"], ws["F2"] = "N", f"=COUNT({ar})"
    ws["F1"].font = BLUE; ws["F1"].fill = YELLOW
    ws["F1"].comment = Comment("Hyperparameter (step 2 of the Python). Change it and watch the table update.", "Claude")

    for j, h in enumerate(["epoch", "m", "b", "dm", "db", "loss (MSE)"]):
        ws.cell(1, 8 + j, h).font = BOLD
    lastrow = 2 + epochs
    for r in range(2, lastrow + 1):
        if r == 2:
            ws["H2"], ws["I2"], ws["J2"] = 0, 0.0, 0.0
            for c in "IJ": ws[f"{c}2"].font = BLUE; ws[f"{c}2"].fill = YELLOW
        else:
            ws[f"H{r}"] = f"=H{r-1}+1"
            ws[f"I{r}"] = f"=I{r-1}-$F$1*K{r-1}"
            ws[f"J{r}"] = f"=J{r-1}-$F$1*L{r-1}"
        err = f"({yr}-I{r}*{xr}-J{r})"
        ws[f"K{r}"] = f"=SUMPRODUCT({err}*-{xr})/$F$2"
        ws[f"L{r}"] = f"=SUMPRODUCT({err}*-1)/$F$2"
        ws[f"M{r}"] = f"=SUMPRODUCT({err}^2)/$F$2"
    ws["I2"].comment = Comment("Initial guess for m (step 3).", "Claude")
    ws["J2"].comment = Comment("Initial guess for b (step 3).", "Claude")

    ws["E4"] = "Final m (fit units)"; ws["F4"] = f"=I{lastrow}"
    ws["E5"] = "Final b (fit units)"; ws["F5"] = f"=J{lastrow}"
    if mode == "raw":
        ws["E7"] = "m in original x units"; ws["F7"] = "=F4"
        ws["E8"] = "b in original x units"; ws["F8"] = "=F5"
    elif mode == "center":
        ws["E7"] = "m in original x units"; ws["F7"] = "=F4"
        ws["E8"] = "b in original x units"; ws["F8"] = f"=F5-F4*AVERAGE({ar})"
    else:
        ws["E7"] = "m in original x units"; ws["F7"] = f"=F4/STDEVP({ar})"
        ws["E8"] = "b in original x units"; ws["F8"] = f"=F5-F4*AVERAGE({ar})/STDEVP({ar})"
    for c in ("F7", "F8"): ws[c].font = BOLD
    ws["E10"] = "Check: SLOPE()"; ws["F10"] = f"=SLOPE($B$2:$B${last},{ar})"
    ws["E11"] = "Check: INTERCEPT()"; ws["F11"] = f"=INTERCEPT($B$2:$B${last},{ar})"
    ws["E13"] = "Yellow cells = inputs you edit. Each row of H:M is one loop iteration."
    if note: ws["E14"] = note

    for row in ws.iter_rows():
        for c in row:
            if c.font.name != "Arial":
                c.font = Font(name="Arial", bold=c.font.bold, color=c.font.color)
    for k, w in {"A": 10, "B": 10, "C": 14, "E": 24, "F": 12, "H": 8, "I": 12, "J": 12, "K": 12, "L": 12, "M": 12}.items():
        ws.column_dimensions[k].width = w

    ch = LineChart(); ch.title = "Loss vs epoch"; ch.y_axis.title = "MSE"; ch.x_axis.title = "epoch"
    ch.add_data(Reference(ws, min_col=13, min_row=1, max_row=lastrow), titles_from_data=True)
    ch.set_categories(Reference(ws, min_col=8, min_row=2, max_row=lastrow))
    ch.legend = None; ch.height, ch.width = 7, 12
    ws.add_chart(ch, "E16")

    if mode != "raw" or n > 4:
        # data + fitted line (fit uses the original-units m and b)
        ws["D1"] = "fit"
        for i in range(2, last + 1): ws[f"D{i}"] = f"=$F$7*A{i}+$F$8"
        sc = ScatterChart(); sc.title = "Data and fitted line"; sc.x_axis.title = xname; sc.y_axis.title = yname
        xs = Reference(ws, min_col=1, min_row=2, max_row=last)
        s1 = Series(Reference(ws, min_col=2, min_row=1, max_row=last), xs, title_from_data=True)
        s1.marker.symbol = "circle"; s1.graphicalProperties.line.noFill = True
        s2 = Series(Reference(ws, min_col=4, min_row=1, max_row=last), xs, title_from_data=True)
        s2.marker.symbol = "none"
        sc.series += [s1, s2]; sc.height, sc.width = 7, 12
        ws.add_chart(sc, "E31")

wb = Workbook()
ws = wb.active; ws.title = "Example raw x"
build(ws, [1, 2, 3, 4], [2, 4, 6, 8], "raw", 0.1, 1000)
build(wb.create_sheet("Example centered x"), [1, 2, 3, 4], [2, 4, 6, 8], "center", 0.5, 50)
wx, wy = [p[0] for p in wall], [p[1] for p in wall]
build(wb.create_sheet("Wall raw x"), wx, wy, "raw", 0.0004, 1000, "Reflection", "BothSpeed",
      "Raw reflection (6-96) forces a tiny learning rate: 0.0005 already blows up, and 0.0004 is still far off after 1000 rows.")
build(wb.create_sheet("Wall standardized x"), wx, wy, "standardize", 0.5, 50, "Reflection", "BothSpeed",
      "Data: annacode-data.csv rows with Reflection > 0 and BothSpeed != 0 (robot sees the wall and is still slowing), sorted by Reflection.")
wb.save(HERE / "gradient_descent.xlsx")
print("wrote", HERE / "gradient_descent.xlsx")

# chart refs sanity check
import zipfile, re
z = zipfile.ZipFile(HERE / "gradient_descent.xlsx")
print(sorted({m for nme in z.namelist() if "charts/chart" in nme for m in re.findall(r"<f>'([^']+)'", z.read(nme).decode())}))
print("wall points", len(wall))
